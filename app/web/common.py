"""Consultas e utilidades compartilhadas pelas telas do admin e do tenant (painel, séries, uploads, downloads)."""
from __future__ import annotations

import csv
import gzip
import io
import json
import os
import zipfile
from datetime import datetime, timedelta, timezone

from flask import Response, current_app, stream_with_context
from sqlalchemy import func, insert, select, text

from ..db import Session, utcnow
from ..engine import formats
from ..models import AllowedIP, Destination, Event, EventDelivery, Parser, Source, StudioJob, Tenant, Upload, User
from ..services import catalog, settings
from ..services.changes import ServiceError

TEXT_EXT = (".txt", ".log", ".json", ".jsonl", ".ndjson", ".csv", ".syslog", ".cef", ".leef", ".gz", ".xlsx")


def nav_counts() -> dict:
    try:
        unp = Session.execute(text("select count(*) from events where status = 'unparsed'")).scalar() or 0
        derr = Session.execute(select(func.count()).select_from(Destination).where(Destination.status == "err", Destination.active.is_(True))).scalar()
        studio = Session.execute(select(func.count()).select_from(StudioJob).where(StudioJob.status == "done")).scalar()
        return {"nav_counts": {"unparsed": unp, "dest_err": derr, "studio": studio}}
    except Exception:  # noqa: BLE001
        Session.rollback()
        return {"nav_counts": {}}


def syslog_info() -> dict:
    c = current_app.config
    ca_note = ("Certificado público (Let's Encrypt, cadeia ISRG Root X1). rsyslog e syslog-ng exigem apontar o arquivo de CA do sistema "
               "(ex.: /etc/ssl/certs/ca-certificates.crt); Firebox não envia syslog com TLS — use TCP ou UDP.")
    return {"host": c["SYSLOG_PUBLIC_HOST"], "ip": c["SYSLOG_PUBLIC_IP"], "tls": c["SYSLOG_TLS_PORT"], "tcp": c["SYSLOG_TCP_PORT"],
            "udp": c["SYSLOG_UDP_PORT"], "ca_note": ca_note}


def tenants() -> list[Tenant]:
    return Session.execute(select(Tenant).order_by(Tenant.name)).scalars().all()


def input_parsers() -> list[Parser]:
    return Session.execute(select(Parser).where(Parser.kind == "input").order_by(Parser.vendor, Parser.name)).scalars().all()


def all_sources() -> list[Source]:
    return Session.execute(select(Source).join(Tenant, Tenant.id == Source.tenant_id).order_by(Tenant.name, Source.name)).scalars().all()


def sources_of(tid) -> list[Source]:
    return Session.execute(select(Source).where(Source.tenant_id == tid).order_by(Source.name)).scalars().all()


def source_names() -> dict:
    return {i: f"{t} · {n}" for i, n, t in Session.execute(select(Source.id, Source.name, Tenant.name).join(Tenant, Tenant.id == Source.tenant_id))}


def tenant_stats(ids: list[int]) -> dict:
    out = {i: {"sources": 0, "dests": 0, "ips": 0, "users": 0, "last": None} for i in ids}
    for tid, n, last in Session.execute(select(Source.tenant_id, func.count(), func.max(Source.last_event_at)).group_by(Source.tenant_id)):
        if tid in out:
            out[tid]["sources"], out[tid]["last"] = n, last
    for tid, n in Session.execute(select(Destination.tenant_id, func.count()).group_by(Destination.tenant_id)):
        if tid in out:
            out[tid]["dests"] = n
    for tid, n in Session.execute(select(AllowedIP.tenant_id, func.count()).where(AllowedIP.active.is_(True)).group_by(AllowedIP.tenant_id)):
        if tid in out:
            out[tid]["ips"] = n
    for tid, n in Session.execute(select(User.tenant_id, func.count()).where(User.tenant_id.is_not(None)).group_by(User.tenant_id)):
        if tid in out:
            out[tid]["users"] = n
    return out


def fw_status() -> dict:
    from ..services import sources as ssvc
    import hashlib
    desired = ssvc.firewall_state()
    h = hashlib.sha256(json.dumps(desired, sort_keys=True).encode()).hexdigest()[:16]
    st = settings.get("fw_status") or {}
    return {"desired": desired, "hash": h, "applied": st, "in_sync": st.get("hash") == h and st.get("ok", False)}


def _series(kind: str, ref_id: int | None, minutes: int = 60) -> list[dict]:
    q = """select minute, sum(a) a, sum(b) b, sum(c) c from metrics_minute where kind = :k and minute > now() - make_interval(mins => :m)
           {extra} group by minute order by minute"""
    params = {"k": kind, "m": minutes}
    extra = ""
    if ref_id is not None:
        extra, params["r"] = "and ref_id = :r", ref_id
    rows = Session.execute(text(q.format(extra=extra)), params).all()
    by = {r.minute.replace(tzinfo=timezone.utc) if r.minute.tzinfo is None else r.minute: r for r in rows}
    now = utcnow().replace(second=0, microsecond=0)
    out = []
    for i in range(minutes - 1, -1, -1):
        m = now - timedelta(minutes=i)
        r = by.get(m)
        out.append({"m": m, "a": int(r.a) if r else 0, "b": int(r.b) if r else 0, "c": int(r.c) if r else 0})
    return out


def source_series(sid):
    return _series("s", sid)


def dest_series(did):
    return _series("d", did)


def dashboard_data(tenant_id: int | None) -> dict:
    tf = "and tenant_id = :t" if tenant_id else ""
    p = {"t": tenant_id}
    k = Session.execute(text(f"""select count(*) filter (where received_at > now() - interval '5 minutes') last5,
                                        count(*) filter (where received_at > now() - interval '60 minutes') last60,
                                        count(*) filter (where status = 'parsed' and received_at > now() - interval '60 minutes') parsed60,
                                        count(*) filter (where status = 'unparsed' and received_at > now() - interval '60 minutes') unparsed60,
                                        count(*) filter (where status = 'unparsed') unparsed_all,
                                        count(*) filter (where status = 'received') queued,
                                        count(*) total
                                 from events where true {tf}"""), p).mappings().one()
    pend = Session.execute(text(f"""select count(*) from event_deliveries d join events e on e.id = d.event_id
                                    where d.status = 'pending' {tf.replace('tenant_id', 'e.tenant_id')}"""), p).scalar() or 0
    dq = select(Destination).order_by(Destination.status.desc(), Destination.name)
    sq = select(Source).order_by(Source.last_event_at.desc().nulls_last())
    if tenant_id:
        dq, sq = dq.where(Destination.tenant_id == tenant_id), sq.where(Source.tenant_id == tenant_id)
    dests = Session.execute(dq).scalars().all()
    srcs = Session.execute(sq.limit(12)).scalars().all()
    series = []
    if tenant_id:
        ids = [s.id for s in Session.execute(select(Source.id).where(Source.tenant_id == tenant_id)).scalars()] if False else \
              [i for (i,) in Session.execute(select(Source.id).where(Source.tenant_id == tenant_id))]
        agg = {}
        for sid in ids:
            for pt in _series("s", sid):
                a = agg.setdefault(pt["m"], {"m": pt["m"], "a": 0, "b": 0, "c": 0})
                a["a"] += pt["a"]
                a["b"] += pt["b"]
                a["c"] += pt["c"]
        series = sorted(agg.values(), key=lambda x: x["m"]) or _series("s", -1)
    else:
        series = _series("s", None)
    hb = settings.get("worker_heartbeat") or {}
    ing = settings.get("ingest_status") or {}
    return {"k": dict(k), "pending": pend, "dests": dests, "srcs": srcs, "series": series, "eps": round((k["last5"] or 0) / 300.0, 2),
            "worker": hb, "ingest": ing, "fw": fw_status() if not tenant_id else None, "syslog": syslog_info(),
            "parsed_pct": round(100.0 * k["parsed60"] / (k["parsed60"] + k["unparsed60"]), 1) if (k["parsed60"] + k["unparsed60"]) else None}


# ------------------------------------------------------------------------------------------------ eventos
def events_query(args, tenant_id: int | None = None):
    q = select(Event)
    filt = {}
    tid = tenant_id or (int(args["tenant_id"]) if args.get("tenant_id") else None)
    if tid:
        q, filt["tenant_id"] = q.where(Event.tenant_id == tid), tid
    if args.get("source_id"):
        q, filt["source_id"] = q.where(Event.source_id == int(args["source_id"])), int(args["source_id"])
    st = args.get("status")
    if st in ("parsed", "unparsed", "received"):
        q, filt["status"] = q.where(Event.status == st), st
    mins = int(args.get("minutes") or 0)
    if mins:
        q, filt["minutes"] = q.where(Event.received_at > utcnow() - timedelta(minutes=min(mins, 1440))), mins
    if args.get("q"):
        q, filt["q"] = q.where(Event.raw.ilike(f"%{args['q'][:200]}%")), args["q"][:200]
    return q, filt


def download_response(q, fmt: str, basename: str, include_unparsed: bool = False):
    """Arquivo para download no formato escolhido, gerado em streaming (uma linha por evento)."""
    choices = dict(catalog.output_choices())
    if fmt not in choices:
        fmt = "ocsf_json"
    spec = catalog.output_spec(fmt)
    ext = formats.FILE_EXT.get(fmt, "txt" if spec is None or spec.get("serializer") not in ("json",) else "jsonl")
    tnames = {t.id: t.name for t in tenants()}
    snames = dict(Session.execute(select(Source.id, Source.name)).all())

    def gen():
        if fmt == "csv":
            yield formats.csv_header() + "\n"
        for ev in Session.execute(q.execution_options(yield_per=500)).scalars():
            if ev.status != "parsed" and not include_unparsed and fmt != "raw":
                continue
            try:
                out = formats.render(fmt, ev.event if ev.status == "parsed" else None, ev.raw,
                                     {"tenant": tnames.get(ev.tenant_id), "source": snames.get(ev.source_id)}, spec)
                yield formats.to_line(out) + "\n"
            except Exception as e:  # noqa: BLE001
                yield formats.to_line({"erro": str(e)[:200], "raw": ev.raw}) + "\n"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    return Response(stream_with_context(gen()), mimetype="text/plain; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{basename}-{fmt}-{stamp}.{ext}"'})


# ------------------------------------------------------------------------------------------------ upload
def _xlsx_lines(data: bytes) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip().lower() for c in next(rows, [])]
    col = None
    for name in ("raw log", "raw_log", "raw", "log", "message", "mensagem", "linha"):
        if name in header:
            col = header.index(name)
            break
    out = []
    if col is None:  # sem cabeçalho reconhecido: usa a coluna de texto mais longa da 1ª linha de dados
        first = list(header)
        col = max(range(len(first)), key=lambda i: len(first[i])) if first else 0
        out.append(str(next(iter([first[col]]), "")))
    for r in rows:
        if r and len(r) > col and r[col] not in (None, ""):
            out.append(str(r[col]))
    return out


def read_upload_lines(f) -> list[str]:
    name = (f.filename or "").lower()
    data = f.read()
    if name.endswith(".gz"):
        try:
            data = gzip.decompress(data)
        except OSError as e:
            raise ServiceError("Arquivo .gz inválido.") from e
        name = name[:-3]
    if name.endswith(".xlsx"):
        try:
            return _xlsx_lines(data)
        except (zipfile.BadZipFile, KeyError, StopIteration, ValueError) as e:
            raise ServiceError("Planilha .xlsx inválida.") from e
    if name.endswith(".csv"):
        txt = data.decode("utf-8-sig", "replace")
        rows = list(csv.reader(io.StringIO(txt)))
        if rows and any(h.strip().lower() in ("raw log", "raw_log", "raw") for h in rows[0]):
            col = [h.strip().lower() for h in rows[0]].index(next(h.strip().lower() for h in rows[0] if h.strip().lower() in ("raw log", "raw_log", "raw")))
            return [r[col] for r in rows[1:] if len(r) > col and r[col].strip()]
        return [ln for ln in txt.splitlines() if ln.strip()]
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        txt = data.decode("utf-16", "replace")
    else:
        txt = data.decode("utf-8-sig", "replace")
    lines = [ln for ln in txt.splitlines() if ln.strip()]
    if name.endswith(".json") and len(lines) > 1 and txt.lstrip().startswith("["):  # lista JSON: um objeto por evento
        try:
            arr = json.loads(txt)
            if isinstance(arr, list):
                return [json.dumps(x, ensure_ascii=False, separators=(",", ":")) for x in arr]
        except ValueError:
            pass
    return lines


def read_doc_file(f) -> str:
    name = (f.filename or "").lower()
    data = f.read()[:5_000_000]
    if name.endswith(".pdf"):
        return "(PDF anexado: cole o texto relevante no campo de documentação — leitura de PDF não habilitada)"
    return data.decode("utf-8", "replace")[:500_000]


def ingest_upload(t: Tenant, f, *, source_id=None, parser_id=None, send=False, by="") -> Upload:
    if f is None or not f.filename:
        raise ServiceError("Escolha o arquivo.")
    if not f.filename.lower().endswith(TEXT_EXT):
        raise ServiceError("Formato não aceito. Envie .txt, .log, .json, .jsonl, .csv, .xlsx ou .gz.")
    if source_id:
        s = Session.get(Source, int(source_id))
        if s is None or s.tenant_id != t.id:
            raise ServiceError("Fonte inválida para este tenant.")
    if parser_id:
        p = Session.get(Parser, int(parser_id))
        if p is None or p.kind != "input":
            raise ServiceError("Parser inválido.")
    lines = read_upload_lines(f)
    if not lines:
        raise ServiceError("O arquivo não tem linhas.")
    if len(lines) > 500_000:
        raise ServiceError("Arquivo grande demais (máximo de 500 mil linhas por envio).")
    hours = float(settings.get("upload_retention_hours") or 2)
    up = Upload(tenant_id=t.id, source_id=int(source_id) if source_id else None, parser_id=int(parser_id) if parser_id else None,
                filename=os.path.basename(f.filename)[:255], size=sum(len(x) for x in lines), lines=len(lines),
                send_to_destinations=bool(send), status="processing", created_by=by, expires_at=utcnow() + timedelta(hours=hours))
    Session.add(up)
    Session.flush()
    now = utcnow()
    rows = [{"tenant_id": t.id, "source_id": up.source_id, "upload_id": up.id, "received_at": now, "peer_ip": "", "transport": "upload",
             "line_no": i + 1, "raw": ln[:262144], "status": "received", "error": "", "pending_deliveries": 0} for i, ln in enumerate(lines)]
    for i in range(0, len(rows), 2000):
        Session.execute(insert(Event), rows[i:i + 2000])
    return up
